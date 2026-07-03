"""Generate Excel file with all Kling AI video generation failure types."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = r"D:\0 北交\0科研-学习myself\场景生成2026.5.25\kling-视频\可灵视频失败类型汇总.xlsx"

# ── Styles ──
header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

cat_font = Font(name="微软雅黑", size=10, bold=True, color="1F4E79")
cat_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

body_font = Font(name="微软雅黑", size=10)
body_align = Alignment(vertical="center", wrap_text=True)

# Category header styles
cat_submitted_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # orange tint
cat_server_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # yellow tint
cat_quality_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # green tint

thin_border = Border(
    left=Side(style="thin", color="C0C0C0"),
    right=Side(style="thin", color="C0C0C0"),
    top=Side(style="thin", color="C0C0C0"),
    bottom=Side(style="thin", color="C0C0C0"),
)

# ── Data ──
headers = ["序号", "所属阶段", "失败类别", "失败类型", "具体表现", "可能原因", "建议操作"]

# Phase 1: Submission errors
phase1_data = [
    ["提交阶段", "配置缺失", "API Key 未配置", "提示「请先在设置中配置可灵 API Keys」", "未设置 Access Key / Secret Key", "在设置对话框填写正确的 API 密钥"],
    ["提交阶段", "鉴权失败", "鉴权失败 (401)", "可灵 API 返回 401 Unauthorized", "AK/SK 错误、过期、或被吊销", "检查并重新生成 API Key，确认 AK/SK 填写无误"],
    ["提交阶段", "鉴权失败", "百炼 API Key 无效", "DashScope 返回认证错误", "百炼 API Key 填写错误或未开通服务", "检查百炼控制台 API Key，确认视频生成服务已开通"],
    ["提交阶段", "图片文件", "图片文件不存在", "「图片文件不存在: .../xxx.png」", "本地 /uploads/ 路径下的图片被删除或未上传", "重新上传参考图后再试"],
    ["提交阶段", "图片格式", "图片格式无效 (base64)", "「图片格式无效: Kling API 需要 base64 编码」", "传入的图片数据不是合法的 base64 或 URL", "确保图片为 base64 data URI 或远程 URL 格式"],
    ["提交阶段", "图片格式", "尾帧图片格式无效", "首尾帧模式下尾帧 base64 不合法", "尾帧图片路径无效或已损坏", "检查尾帧图片是否存在，重新上传"],
    ["提交阶段", "图片解析", "图片无法解析", "「无法解析图片: ...」", "图片不是 data URI / 远程URL / 本地路径中任一格式", "使用上传功能上传图片，或提供有效的远程 URL"],
    ["提交阶段", "API 错误", "Kling 服务端通用错误", "「Kling API error (xxx): ...」", "Kling 服务返回非 200 状态码，可能服务繁忙", "稍等片刻重试；若持续失败联系可灵技术支持"],
    ["提交阶段", "API 错误", "DashScope 提交失败", "「DashScope error (xxx): ...」", "阿里云百炼返回错误，可能服务不可用或参数不兼容", "检查参数是否匹配所选模型，或切换 provider"],
]

# Phase 2: Server-side task failures (task_status_msg patterns)
phase2_data = [
    # 内容审核
    ["服务端任务", "内容审核", "提示词违规", "提示词包含敏感词或违规内容，任务被拒绝", "安全审核机制拦截了提示词中的敏感词", "修改提示词，避免政治、暴力、色情等敏感内容"],
    ["服务端任务", "内容审核", "图片内容违规", "参考图片未通过内容安全审核", "上传的图片包含违规元素", "更换参考图，确保图片内容合规"],
    ["服务端任务", "内容审核", "生成结果违规", "生成的视频内容被安全审核拦截", "模型输出了不符合安全规范的内容", "调整提示词方向，避免高风险内容描述"],
    ["服务端任务", "内容审核", "人脸/版权检测失败", "检测到未授权人脸或疑似版权内容", "参考图包含知名人物面部或版权素材", "使用原创/授权素材，或遮挡面部特征"],
    # 技术/参数
    ["服务端任务", "技术参数", "模型推理失败", "模型内部生成过程出错", "推理服务异常或输入参数触发边缘 case", "更换模型版本或降低 cfg_scale 尝试"],
    ["服务端任务", "技术参数", "参数不支持", "所选参数组合不被当前模型支持", "模型的 mode/duration/resolution 组合无效", "查看模型配置表，选择支持的参数组合"],
    ["服务端任务", "技术参数", "参考图质量不足", "参考图清晰度不够或特征不明显", "图片分辨率过低或主体不清晰", "使用高清、主体突出的参考图（建议 ≥720p）"],
    ["服务端任务", "技术参数", "首尾帧不匹配", "首帧和尾帧差异过大，无法生成连贯视频", "首尾帧构图、内容差异超出模型处理能力", "选择内容连贯的首尾帧图片，确保过渡可行性"],
    ["服务端任务", "技术参数", "主体识别失败", "指定的 element/subject 未被识别", "element_id 无效或主体已被删除", "重新创建主体，确认 element_id 正确"],
    ["服务端任务", "技术参数", "视频参考分析失败", "reference video 无法被解析或分析", "参考视频格式不支持或时长超出限制", "使用标准格式（MP4），控制视频时长和大小"],
    # 资源/配额
    ["服务端任务", "资源配额", "配额/余额不足", "API 调用次数用完或账户余额不足", "资源包耗尽、免费额度用完、欠费", "充值或购买资源包，检查账户余额"],
    ["服务端任务", "资源配额", "并发/频率限制", "请求过于频繁被限流", "短时间内提交了过多任务", "降低请求频率，间隔 500ms 以上；使用队列"],
    ["服务端任务", "资源配额", "模型/服务不可用", "所选模型临时下线或服务维护中", "Kling/百炼服务升级或故障", "切换备用模型，等待服务恢复（通常几分钟）"],
    ["服务端任务", "资源配额", "任务超时", "任务处理超过最大时限", "视频复杂度过高或服务排队拥堵", "简化提示词/缩短时长，非高峰时段重试"],
    # 其他
    ["服务端任务", "其他", "未知服务端错误", "返回 generic error / internal server error", "服务端未分类的异常", "记录 task_id 和错误信息，联系技术支持"],
    ["服务端任务", "其他", "回调/网络错误", "callback / network error", "网络波动导致任务结果回调失败", "检查网络连接，手动轮询任务状态"],
]

# Phase 3: Video content quality issues (our new reject reasons)
phase3_data = [
    ["视频内容质量", "画面崩坏", "画面撕裂/花屏", "出现大面积色块、像素错乱、马赛克", "模型推理异常、分辨率过高", "降低分辨率 (如 4K→1080P)，切换模型版本重试"],
    ["视频内容质量", "画面崩坏", "黑屏/白屏", "全片或部分片段纯黑/纯白", "参考图未被正确理解、首尾帧冲突", "检查参考图是否有效，尝试不加参考图生成"],
    ["视频内容质量", "画面崩坏", "严重畸变", "人脸/身体拉伸扭曲到完全不可辨认", "提示词描述矛盾、cfg_scale 过高", "降低 cfg_scale，简化提示词中的约束"],
    ["视频内容质量", "画面崩坏", "帧冻结", "画面完全静止不动", "图生视频时参考图 motion 未被驱动", "更换参考图，增加描述动作的提示词内容"],
    ["视频内容质量", "主体问题", "主体变形", "手脚扭曲、多出肢体、面部崩塌、五官错位", "模型对复杂人体结构理解不足", "简化动作描述，降低 cfg_scale，使用主体绑定功能"],
    ["视频内容质量", "主体问题", "主体丢失", "生成中途主体突然消失或变成别的物体", "提示词未持续强调主体存在", "在提示词中持续描述主体，设定更短的时长"],
    ["视频内容质量", "主体问题", "主体漂移", "人物/物体位置不断偏移，甚至飘出画外", "缺少空间锚点描述", "添加场景/背景描述固定空间关系"],
    ["视频内容质量", "主体问题", "主体克隆", "同一个物体/人物在画面中出现多个副本", "模型误解提示词中的数量描述", "明确使用「一个人」「单个」等数量限定词"],
    ["视频内容质量", "主体问题", "主体融合", "两个物体/人物诡异融合在一起", "多个主体距离过近，边界模糊", "拉开主体间距描述，使用 element 绑定区分"],
    ["视频内容质量", "运动异常", "鬼畜/抖动", "画面高频抖动、抽搐、像跳帧", "运动幅度参数与内容不匹配", "降低 cfg_scale，使用更稳定的模型版本"],
    ["视频内容质量", "运动异常", "动作不自然", "走路像滑冰、手势诡异、运动违反物理规律", "模型对物理规律学习不足", "简化动作描述，使用更高级模型（如 v3）"],
    ["视频内容质量", "运动异常", "瞬移", "物体突然从一处跳到另一处，无过渡", "帧间一致性不足", "降低时长，减少场景变化幅度"],
    ["视频内容质量", "运动异常", "逆时序", "动作倒放、水往上流等违反因果现象", "模型对因果/时序关系理解偏差", "强调正向时序关键词，如「逐渐」「慢慢」"],
    ["视频内容质量", "运动异常", "闪烁", "画面亮度/色彩在帧间剧烈波动", "模型输出不稳定", "使用 pro 模式，提升生成质量"],
    ["视频内容质量", "内容偏差", "完全跑偏", "生成内容跟提示词完全无关", "提示词被模型误解析或忽略", "简化提示词，使用更直接的关键词描述"],
    ["视频内容质量", "内容偏差", "部分遗漏", "提示词中某些关键元素被忽略", "提示词过长或元素间有冲突", "缩短提示词，减少同时要求的元素数量"],
    ["视频内容质量", "内容偏差", "风格错误", "要求写实出了动画风、要求3D出了2D", "风格关键词权重不够", "加强风格描述，添加负面提示词排除错误风格"],
    ["视频内容质量", "内容偏差", "构图失败", "主体被裁切、取景角度怪异、比例失调", "缺少构图/镜头指令", "添加镜头指令（如 wide shot, close-up, centered）"],
    ["视频内容质量", "内容偏差", "文字幻觉", "画面中出现奇怪的文字/logo/水印", "模型训练数据中的文字被错误再现", "添加负面提示词「文字, logo, watermark」，降低 cfg_scale"],
    ["视频内容质量", "内容偏差", "首尾帧断裂", "首尾帧模式下中间过渡生硬跳跃，前后风格不一致", "首尾帧差异过大或模型版本不支持", "选择差异较小的首尾帧，使用支持首尾帧的模型"],
    ["视频内容质量", "技术质量", "清晰度不足", "模糊、涂抹感重、像低码率视频", "std 模式画质限制或参考图质量低", "使用 pro/4k 模式，提高参考图分辨率"],
    ["视频内容质量", "技术质量", "色彩异常", "偏色严重、饱和度异常、死黑/过曝", "模型色彩理解偏差或光照描述冲突", "添加色彩/光照描述：natural lighting, balanced color"],
    ["视频内容质量", "技术质量", "多镜头割裂", "Omni 多镜头模式下不同镜头间风格光照突然变化", "各镜头独立生成缺乏全局一致性约束", "统一各镜头的光照/色彩描述，使用 Omni 模型"],
    ["视频内容质量", "技术质量", "音画不同步", "生成带音频的视频时声音和画面不对应", "音频驱动的口型/动作匹配失败", "确保音频和视频描述内容一致，使用支持音频的模型"],
]


def write_sheet(ws, data, title_fill=None):
    """Write headers and data rows to worksheet."""
    # Headers
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    seq = 0
    for row_data in data:
        seq += 1
        row_num = ws.max_row + 1

        phase = row_data[0]
        if phase == "提交阶段":
            row_fill = cat_submitted_fill
        elif phase == "服务端任务":
            row_fill = cat_server_fill
        else:
            row_fill = cat_quality_fill

        values = [seq, *row_data]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = body_font
            cell.alignment = body_align
            cell.fill = row_fill
            cell.border = thin_border

    # Set column widths
    col_widths = [6, 14, 16, 26, 46, 40, 42]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze header
    ws.freeze_panes = "A2"
    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


wb = openpyxl.Workbook()

# ── Sheet 1: Comprehensive ──
ws1 = wb.active
ws1.title = "失败类型汇总"
write_sheet(ws1, phase1_data + phase2_data + phase3_data)

# ── Sheet 2: 提交阶段 ──
ws2 = wb.create_sheet("1-提交阶段失败")
write_sheet(ws2, phase1_data)

# ── Sheet 3: 服务端任务 ──
ws3 = wb.create_sheet("2-服务端任务失败")
write_sheet(ws3, phase2_data)

# ── Sheet 4: 视频内容质量 ──
ws4 = wb.create_sheet("3-视频内容质量")
write_sheet(ws4, phase3_data)

wb.save(OUTPUT)
print(f"Excel saved to: {OUTPUT}")
print(f"   Sheets: {wb.sheetnames}")
print(f"   Total rows: {ws1.max_row - 1} failure types")
